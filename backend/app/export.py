"""Excel export of the reconciliation for the finance team."""

from __future__ import annotations

import io
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font

ZAR = '"R" #,##0.00'


def build_workbook(summary: dict, exceptions: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Measure", "Value"])
    ws.append(["Receipts in bank", Decimal(summary["receipts_value"])])
    ws.append(["Matched to invoices", Decimal(summary["matched_value"])])
    # A formula, so the reviewer can see how the unreconciled figure is derived.
    ws.append(["Not yet reconciled", "=B2-B3"])
    ws.append(["Exceptions (listed)", "=SUM(Exceptions!D:D)"])
    ws.append(["Check: difference", '=IF(ROUND(B4-B5,2)=0,"OK","INVESTIGATE")'])
    ws.append(["Open invoices", summary["open_invoices"]])
    ws.append(["Open invoice value", Decimal(summary["open_invoice_value"])])
    for row in (2, 3, 4, 5, 8):
        ws[f"B{row}"].number_format = ZAR
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 18

    ex = wb.create_sheet("Exceptions")
    headers = ["Bank ref", "Date", "Description", "Amount", "Invoice", "Rule", "Difference", "Note", "Action"]
    ex.append(headers)
    for e in exceptions:
        ex.append(
            [
                e["bank_transaction_id"],
                e["posted_on"],
                e["description"],
                Decimal(e["amount"]),
                e["invoice_number"],
                e["rule"],
                Decimal(e["difference"]),
                e["note"],
                "",
            ]
        )
    for row in ex.iter_rows(min_row=2):
        row[3].number_format = ZAR
        row[6].number_format = ZAR
    for cell in ws[1] + ex[1]:
        cell.font = Font(bold=True)
    for col, width in zip("ABCDEFGHI", [14, 11, 40, 13, 12, 14, 12, 40, 20], strict=True):
        ex.column_dimensions[col].width = width
    ex.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
