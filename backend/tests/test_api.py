import io
import json
import time
from pathlib import Path

from openpyxl import load_workbook

from app.security import sign

STATEMENT = (Path(__file__).resolve().parents[1] / "sample_data" / "bank_statement.csv").read_bytes()


def upload(client, content=STATEMENT):
    return client.post("/imports/bank-statement", files={"file": ("statement.csv", content, "text/csv")})


def test_import_is_idempotent(client):
    assert upload(client).json() == {"imported": 7, "skipped_duplicates": 0}
    assert upload(client).json() == {"imported": 0, "skipped_duplicates": 7}


def test_bad_row_rejects_whole_file(client):
    bad = b"id,posted_on,description,amount\nX1,2026-10-01,ok,10.00\nX2,not-a-date,bad,5\n"
    response = upload(client, bad)
    assert response.status_code == 422
    assert "row 3" in response.json()["detail"]["errors"][0]
    assert client.get("/summary").json()["receipts_value"] == "0.00"


def test_full_reconciliation_flow(client):
    upload(client)
    run = client.post("/reconcile").json()
    assert run["by_rule"] == {
        "reference": 2,
        "short_payment": 1,
        "amount_date": 1,
        "ambiguous": 1,
        "unmatched": 1,
    }

    exceptions = {e["bank_transaction_id"]: e for e in client.get("/exceptions").json()}
    assert exceptions["FNB-20261002-002"]["difference"] == "-452.25"
    assert len(exceptions) == 4

    # A person confirms the amount/date suggestion for Ridgeview.
    assert client.post("/matches/FNB-20261006-004/confirm").status_code == 200
    # Short payments can't be confirmed as a full match.
    assert client.post("/matches/FNB-20261002-002/confirm").status_code == 409

    summary = client.get("/summary").json()
    assert summary["matched_value"] == "28350.15"  # 7274.90 + 8977.25 + 12098.00
    assert summary["needs_review"] == 3

    # Re-running keeps confirmed matches and doesn't duplicate anything.
    assert client.post("/reconcile").json()["processed"] == 3


def test_export_workbook(client):
    upload(client)
    client.post("/reconcile")
    response = client.get("/export.xlsx")
    wb = load_workbook(io.BytesIO(response.content))
    assert wb["Summary"]["B4"].value == "=B2-B3"
    assert wb["Exceptions"].max_row == 5  # header + 4 exceptions


def webhook(client, payload: dict, secret="test-secret", timestamp=None):
    body = json.dumps(payload).encode()
    header = sign(secret, body, timestamp or int(time.time()))
    return client.post("/webhooks/payments", content=body, headers={"X-Signature": header})


EVENT = {
    "id": "evt_001",
    "type": "payment.received",
    "data": {
        "transaction_id": "PAY-9001",
        "posted_on": "2026-10-12",
        "description": "Card payment INV-1048",
        "amount": "1150.00",
    },
}


def test_webhook_creates_transaction_once(client):
    assert webhook(client, EVENT).json() == {"status": "processed"}
    assert webhook(client, EVENT).json() == {"status": "duplicate"}
    assert client.post("/reconcile").json()["by_rule"] == {"reference": 1}


def test_webhook_rejects_bad_signature_and_replays(client):
    assert webhook(client, EVENT, secret="wrong").status_code == 401
    assert webhook(client, EVENT, timestamp=int(time.time()) - 3600).status_code == 401
    no_header = client.post("/webhooks/payments", content=json.dumps(EVENT))
    assert no_header.status_code == 401
